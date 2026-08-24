"""
Serwis importu pełnego eksportu XLSX z Comarch ERP Optima.

Import jest pełną migawką (nie dopisywaniem wierszy). Proces dwuetapowy:
1) `analyze_workbook` — czysta analiza bez zapisu do bazy, zwraca raport
   walidacji i przygotowane rekordy do podglądu;
2) `commit_batch` — zapisuje wynik analizy transakcyjnie jako nową,
   aktywną migawkę (poprzednia migawka pozostaje w bazie do audytu).
"""

import logging
from dataclasses import asdict, dataclass, field
from decimal import Decimal

import openpyxl
from django.db import transaction
from django.utils import timezone

from pawilony.models import ImportBatch, PavilionSnapshot
from pawilony.services.header_mapping import build_header_map
from pawilony.services.hours import PavilionEquipment, calculate_hours, get_operation_hours_map
from pawilony.services.module_parser import is_counted_module, parse_module, project_key
from pawilony.services.normalization import normalize_bool, normalize_cell
from pawilony.services.status_rules import classify_status

logger = logging.getLogger("pawilony.audit")

SHEET_NAME = "Lista zasobów"

OD_REKI_LABEL = "od ręki"
OD_REKI_ALLOWED_STATUSES = {"Produkcja Zabrze", "Produkcja Czekanów"}


def _is_od_reki(rodzaj_norm) -> bool:
    """
    'Rodzaj' == 'Od ręki' — dotyczy pojedynczej wartości oraz komórek
    wieloznacznych (np. 'Od ręki\\nZamówiony'), żeby nie zaniżać ograniczenia
    przez niejednoznaczność danych źródłowych.
    """
    if rodzaj_norm.is_ambiguous:
        return any(v.strip().lower() == OD_REKI_LABEL for v in rodzaj_norm.distinct_values)
    return (rodzaj_norm.value or "").strip().lower() == OD_REKI_LABEL


class ImportAnalysisError(Exception):
    """Błąd blokujący dalsze przetwarzanie pliku (np. brak arkusza lub kolumn)."""


@dataclass
class ImportReport:
    total_rows: int = 0
    active_count: int = 0
    ended_count: int = 0
    unrecognized_status_count: int = 0
    empty_status_count: int = 0
    module_skipped_count: int = 0
    od_reki_excluded_count: int = 0
    conflict_count: int = 0
    duplicate_count: int = 0
    unrecognized_values: list[dict] = field(default_factory=list)
    conflicts: list[dict] = field(default_factory=list)
    warnings: list[dict] = field(default_factory=list)
    missing_columns: list[str] = field(default_factory=list)
    unrecognized_headers: list[str] = field(default_factory=list)
    sheet_name: str = SHEET_NAME
    counted_hydraulic_hours: str = "0"
    counted_welding_hours: str = "0"
    counted_fibo_wood_hours: str = "0"
    custom_count: int = 0
    standard_count: int = 0


def _cell_value(row: tuple, field_map: dict, field_name: str):
    idx = field_map.get(field_name)
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def analyze_workbook(
    file_obj, exclude_od_reki_before_production: bool = True
) -> tuple[ImportReport, list[dict]]:
    """
    Parsuje i waliduje skoroszyt. Nie zapisuje niczego do bazy.

    `exclude_od_reki_before_production` steruje przełącznikiem z Konfiguracji
    mocy produkcyjnych: gdy True (domyślnie), pawilony 'Od ręki' nie liczą się
    do kolejki, dopóki nie mają statusu Produkcja Zabrze/Czekanów. Gdy False,
    'Od ręki' liczy się tak samo jak 'Zamówiony' na każdym aktywnym statusie.
    """
    try:
        workbook = openpyxl.load_workbook(file_obj, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — chcemy czytelny komunikat, nie surowy traceback
        raise ImportAnalysisError(f"Nie można odczytać pliku XLSX: {exc}") from exc

    if SHEET_NAME not in workbook.sheetnames:
        raise ImportAnalysisError(
            f"Brak wymaganego arkusza '{SHEET_NAME}' w pliku. Dostępne arkusze: {', '.join(workbook.sheetnames)}"
        )

    sheet = workbook[SHEET_NAME]
    rows_iter = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration as exc:
        raise ImportAnalysisError("Arkusz jest pusty — brak wiersza nagłówków.") from exc

    field_map, missing_required, unrecognized_headers = build_header_map(list(header_row))
    report = ImportReport(
        missing_columns=missing_required,
        unrecognized_headers=unrecognized_headers,
    )

    if missing_required:
        raise ImportAnalysisError(
            "Brak wymaganych kolumn w pliku: " + ", ".join(missing_required)
        )

    op_hours = get_operation_hours_map()
    records: list[dict] = []
    seen_active_kods: dict[str, int] = {}

    # Pierwszy przebieg: policz liczbę modułów każdego projektu (wiersze
    # oznaczone tym samym "MODUŁ n" w nazwie, po usunięciu oznaczenia modułu)
    # — statyka/kratownica są liczone razy liczba modułów, więc musimy znać
    # ją przed policzeniem godzin dla wiersza, który faktycznie wejdzie do kolejki.
    all_rows = list(rows_iter)
    module_counts_by_project: dict[str, int] = {}
    for row in all_rows:
        if row is None or all(v is None for v in row):
            continue
        nazwa_val = normalize_cell(_cell_value(row, field_map, "nazwa")).value or ""
        mres = parse_module(nazwa_val)
        if mres.present and not mres.conflict:
            key = project_key(nazwa_val)
            module_counts_by_project[key] = module_counts_by_project.get(key, 0) + 1

    hydraulic_total = Decimal("0")
    welding_total = Decimal("0")
    fibo_wood_total = Decimal("0")
    custom_count = 0
    standard_count = 0

    for row_number, row in enumerate(all_rows, start=2):
        if row is None or all(v is None for v in row):
            continue
        report.total_rows += 1

        raw_values = {
            field_name: _cell_value(row, field_map, field_name) for field_name in field_map
        }

        kod_norm = normalize_cell(raw_values.get("kod"))
        nazwa_norm = normalize_cell(raw_values.get("nazwa"))
        typ_norm = normalize_cell(raw_values.get("typ"))
        stan_norm = normalize_cell(raw_values.get("stan_zasobow"))
        rodzaj_norm = normalize_cell(raw_values.get("rodzaj"))
        status_norm = normalize_cell(raw_values.get("status_proces"))
        termin_norm = normalize_cell(raw_values.get("termin_realiz"))

        kod = kod_norm.value or ""
        nazwa = nazwa_norm.value or ""

        conflict_reasons: list[str] = []
        warnings: list[str] = []

        if kod_norm.had_duplicate:
            warnings.append("Kod: powtórzona identyczna wartość w komórce — znormalizowano.")
        if nazwa_norm.had_duplicate:
            warnings.append("Nazwa: powtórzona identyczna wartość w komórce — znormalizowano.")
        if rodzaj_norm.had_duplicate:
            warnings.append("Rodzaj: powtórzona identyczna wartość w komórce — znormalizowano.")
        if rodzaj_norm.is_ambiguous:
            warnings.append(
                f"Rodzaj: wieloznaczna wartość ({', '.join(rodzaj_norm.distinct_values)}) — pole informacyjne."
            )

        status_result = classify_status(status_norm)
        if status_norm.had_duplicate and status_result.classification != "CONFLICT":
            warnings.append("Status procesu: powtórzona identyczna wartość w komórce — znormalizowano.")
        if status_result.classification == "CONFLICT":
            conflict_reasons.append(status_result.reason)
        elif status_result.classification == "UNRECOGNIZED":
            report.unrecognized_values.append(
                {"row": row_number, "kod": kod, "pole": "status_proces", "wartosc": status_result.canonical}
            )
        elif status_result.classification == "EMPTY":
            report.empty_status_count += 1

        module_result = parse_module(nazwa)
        if module_result.conflict:
            conflict_reasons.append(
                f"Nazwa sugeruje moduł, ale numer jest niejednoznaczny: '{module_result.raw_match or nazwa}'"
            )

        def read_attr(name):
            return normalize_cell(raw_values.get(name))

        kuchnia_norm = read_attr("kuchnia")
        toaleta_norm = read_attr("toaleta")
        lazienka_norm = read_attr("lazienka")
        pelna_statyka_bool, statyka_warn = normalize_bool(raw_values.get("pelna_statyka"))
        kratownica_bool, kratownica_warn = normalize_bool(raw_values.get("kratownica"))
        prysznic_bool, prysznic_warn = normalize_bool(raw_values.get("prysznic"))

        for warn in (statyka_warn, kratownica_warn, prysznic_warn):
            if warn:
                warnings.append(warn)
                report.unrecognized_values.append({"row": row_number, "kod": kod, "pole": "logiczne", "wartosc": warn})

        if pelna_statyka_bool and kratownica_bool and status_result.classification == "ACTIVE":
            # Sprawdzamy to tylko dla aktywnych statusów (Logistyka/Produkcja Zabrze/
            # Produkcja Czekanów) — pawilony o innych statusach (wysłane, odebrane...)
            # nigdy nie trafiłyby do backlogu, więc nie ma sensu zgłaszać tu konfliktu.
            conflict_reasons.append(
                "Pełna konstrukcja/statyka i kratownica nie mogą wystąpić jednocześnie."
            )

        def resolve_variant(norm, allowed: set[str], field_label: str) -> str:
            if norm.is_ambiguous:
                conflict_reasons.append(
                    f"{field_label}: wieloznaczna wartość ({', '.join(norm.distinct_values)})"
                )
                return ""
            if norm.value is None:
                return ""
            if norm.value.strip().capitalize() not in allowed and norm.value.strip() not in allowed:
                report.unrecognized_values.append(
                    {"row": row_number, "kod": kod, "pole": field_label, "wartosc": norm.value}
                )
                warnings.append(f"{field_label}: nierozpoznana wartość '{norm.value}'")
                return ""
            for candidate in allowed:
                if candidate.lower() == norm.value.strip().lower():
                    return candidate
            return ""

        kuchnia_val = resolve_variant(kuchnia_norm, {"Standard", "Lux"}, "Kuchnia")
        toaleta_val = resolve_variant(toaleta_norm, {"Standard", "Komfort", "Premium"}, "Toaleta")
        lazienka_val = resolve_variant(lazienka_norm, {"Standard", "Komfort", "Premium"}, "Łazienka")

        status_classification = status_result.classification
        is_active_status = status_classification == "ACTIVE"
        module_ok = is_counted_module(module_result)
        has_conflict = bool(conflict_reasons)

        # Pawilony "Od ręki" nie wchodzą do kolejki, dopóki nie mają statusu
        # Produkcja Zabrze albo Produkcja Czekanów — samo "Logistyka" (mimo że
        # jest statusem aktywnym dla pawilonów "Zamówiony") ich jeszcze nie liczy.
        od_reki_excluded = (
            exclude_od_reki_before_production
            and is_active_status
            and _is_od_reki(rodzaj_norm)
            and status_result.canonical not in OD_REKI_ALLOWED_STATUSES
        )
        if od_reki_excluded:
            warnings.append(
                "Rodzaj 'Od ręki': pominięto w kolejce do czasu statusu "
                "Produkcja Zabrze/Produkcja Czekanów."
            )
            report.od_reki_excluded_count += 1

        is_counted = is_active_status and module_ok and not has_conflict and not od_reki_excluded

        if is_counted:
            key = kod.lower()
            if key and key in seen_active_kods:
                conflict_reasons.append(
                    f"Duplikat aktywnego kodu — wcześniej wystąpił w wierszu {seen_active_kods[key]}."
                )
                report.duplicate_count += 1
                is_counted = False
            elif key:
                seen_active_kods[key] = row_number

        if status_classification == "ACTIVE" and not module_ok and not module_result.conflict:
            report.module_skipped_count += 1

        if module_result.present and not module_result.conflict:
            row_module_count = module_counts_by_project.get(project_key(nazwa), 1)
        else:
            row_module_count = 1

        equipment = PavilionEquipment(
            kuchnia=kuchnia_val or None,
            toaleta=toaleta_val or None,
            lazienka=lazienka_val or None,
            prysznic=prysznic_bool,
            pelna_statyka=pelna_statyka_bool,
            kratownica=kratownica_bool,
            module_count=row_module_count,
            stolarka_nst=bool(read_attr("stolarka_nst").value),
            zaluzje_fasadowe=bool(read_attr("zaluzje_fasadowe").value),
            rolety=bool(read_attr("rolety").value),
        )
        hours_result = calculate_hours(equipment, op_hours)

        if hours_result.warnings:
            warnings.extend(hours_result.warnings)

        if has_conflict:
            status_classification = "CONFLICT"
            report.conflict_count += 1
            report.conflicts.append(
                {"row": row_number, "kod": kod, "nazwa": nazwa, "powody": conflict_reasons}
            )

        if status_result.classification == "ACTIVE" and not has_conflict:
            report.active_count += 1
        elif status_result.classification == "ENDED":
            report.ended_count += 1
        elif status_result.classification == "UNRECOGNIZED":
            report.unrecognized_status_count += 1

        record_hydraulic = hours_result.hydraulic_hours if is_counted else Decimal("0")
        record_welding = hours_result.welding_hours if is_counted else Decimal("0")
        record_fibo_wood = hours_result.fibo_wood_hours if is_counted else Decimal("0")
        is_custom = is_counted and hours_result.is_custom

        if is_counted:
            hydraulic_total += record_hydraulic
            welding_total += record_welding
            fibo_wood_total += record_fibo_wood
            if is_custom:
                custom_count += 1
            else:
                standard_count += 1

        if warnings:
            report.warnings.append({"row": row_number, "kod": kod, "nazwa": nazwa, "komunikaty": warnings})

        records.append(
            {
                "source_row_number": row_number,
                "kod": kod,
                "nazwa": nazwa,
                "typ": typ_norm.value or "",
                "stan_zasobow": stan_norm.value or "",
                "ilosc_dostepna": str(raw_values.get("ilosc_dostepna") or ""),
                "jm": str(raw_values.get("jm") or ""),
                "oddzial": str(normalize_cell(raw_values.get("oddzial")).value or ""),
                "rodzaj_raw": str(raw_values.get("rodzaj") or ""),
                "rodzaj_normalized": rodzaj_norm.value or "",
                "status_raw": str(raw_values.get("status_proces") or ""),
                "status_normalized": status_result.canonical,
                "status_classification": status_classification,
                "termin_realiz_raw": str(raw_values.get("termin_realiz") or termin_norm.value or ""),
                "module_raw_text": module_result.raw_match,
                "module_number": module_result.number,
                "module_present": module_result.present,
                "module_conflict": module_result.conflict,
                "kuchnia": kuchnia_val,
                "toaleta": toaleta_val,
                "lazienka": lazienka_val,
                "prysznic": prysznic_bool,
                "pelna_statyka": pelna_statyka_bool,
                "kratownica": kratownica_bool,
                "stolarka_nst_raw": str(raw_values.get("stolarka_nst") or ""),
                "zaluzje_fasadowe_raw": str(raw_values.get("zaluzje_fasadowe") or ""),
                "rolety_raw": str(raw_values.get("rolety") or ""),
                "plyty_niestandard_raw": str(raw_values.get("plyty_niestandard") or ""),
                "inne_niestandard_raw": str(raw_values.get("inne_niestandard") or ""),
                "attributes_raw": {k: ("" if v is None else str(v)) for k, v in raw_values.items()},
                "od_reki_excluded": od_reki_excluded,
                "is_counted": is_counted,
                "is_custom": is_custom,
                "hydraulic_hours": str(record_hydraulic),
                "welding_hours": str(record_welding),
                "fibo_wood_hours": str(record_fibo_wood),
                "conflict_reasons": conflict_reasons,
                "warnings": warnings,
            }
        )

    report.counted_hydraulic_hours = str(hydraulic_total)
    report.counted_welding_hours = str(welding_total)
    report.counted_fibo_wood_hours = str(fibo_wood_total)
    report.custom_count = custom_count
    report.standard_count = standard_count

    return report, records


@transaction.atomic
def commit_batch(batch: ImportBatch, records: list[dict], user) -> ImportBatch:
    """Zapisuje przeanalizowane rekordy jako nową aktywną migawkę (transakcyjnie)."""
    snapshots = [
        PavilionSnapshot(
            import_batch=batch,
            kod=r["kod"][:100],
            nazwa=r["nazwa"][:500],
            typ=r["typ"][:50],
            stan_zasobow=r["stan_zasobow"][:200],
            ilosc_dostepna=r["ilosc_dostepna"][:50],
            jm=r["jm"][:20],
            oddzial=r["oddzial"][:200],
            rodzaj_raw=r["rodzaj_raw"][:200],
            rodzaj_normalized=r["rodzaj_normalized"][:200],
            status_raw=r["status_raw"][:300],
            status_normalized=r["status_normalized"][:200],
            status_classification=r["status_classification"],
            termin_realiz_raw=r["termin_realiz_raw"][:200],
            module_raw_text=r["module_raw_text"][:500],
            module_number=r["module_number"],
            module_present=r["module_present"],
            module_conflict=r["module_conflict"],
            kuchnia=r["kuchnia"],
            toaleta=r["toaleta"],
            lazienka=r["lazienka"],
            prysznic=r["prysznic"],
            pelna_statyka=r["pelna_statyka"],
            kratownica=r["kratownica"],
            stolarka_nst_raw=r["stolarka_nst_raw"][:50],
            zaluzje_fasadowe_raw=r["zaluzje_fasadowe_raw"][:50],
            rolety_raw=r["rolety_raw"][:50],
            plyty_niestandard_raw=r["plyty_niestandard_raw"][:200],
            inne_niestandard_raw=r["inne_niestandard_raw"][:200],
            attributes_raw=r["attributes_raw"],
            is_counted=r["is_counted"],
            is_custom=r["is_custom"],
            hydraulic_hours=Decimal(r["hydraulic_hours"]),
            welding_hours=Decimal(r["welding_hours"]),
            fibo_wood_hours=Decimal(r["fibo_wood_hours"]),
            conflict_reasons=r["conflict_reasons"],
            warnings=r["warnings"],
            source_row_number=r["source_row_number"],
        )
        for r in records
    ]
    PavilionSnapshot.objects.bulk_create(snapshots, batch_size=500)

    ImportBatch.objects.filter(is_active_snapshot=True).update(is_active_snapshot=False)
    batch.status = ImportBatch.Status.COMMITTED
    batch.is_active_snapshot = True
    batch.committed_at = timezone.now()
    batch.save(update_fields=["status", "is_active_snapshot", "committed_at"])

    logger.info(
        "Import zatwierdzony: batch_id=%s użytkownik=%s wierszy=%s aktywnych=%s",
        batch.pk,
        getattr(user, "username", "?"),
        len(records),
        sum(1 for r in records if r["is_counted"]),
    )
    return batch


def report_to_dict(report: ImportReport) -> dict:
    return asdict(report)
