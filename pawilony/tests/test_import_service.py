import io
from decimal import Decimal

import openpyxl
import pytest

from pawilony.models import ImportBatch, PavilionSnapshot
from pawilony.services.import_service import ImportAnalysisError, analyze_workbook, commit_batch

HEADER = [
    "Kod", "Nazwa", "Typ", "Stan zasobów", "Ilość dostępna", "Jm",
    "01. ODDZIAŁ (Atrybut)", "02. RODZAJ (Atrybut)", "03. STATUS PROCESU (Atrybut)",
    "07. TERMIN REALIZ. (Atrybut)", "23. PŁYTY NIESTANDAR (Atrybut)",
    "25. PEŁNA/STATYKA (Atrybut)", "24. KRATOWNICA (Atrybut)", "26. KUCHNIA (Atrybut)",
    "28. PRYSZNIC (Atrybut)", "27. TOALETA (Atrybut)", "30. STOLARKA NST (Atrybut)",
    "32. ŻALUZJE FASADOWE (Atrybut)", "31. ROLETY (Atrybut)", "34. ŁAZIENKA (Atrybut)",
    "33. INNE NIESTANDARD (Atrybut)",
]


def _make_row(kod="P001", nazwa="Pawilon 5x3 nr projektu 1/1/2024", typ="TZ", stan="Brak towaru",
              ilosc=1, jm="szt", oddzial="Zabrze", rodzaj="Zamówiony", status="Logistyka",
              termin="", plyty="", statyka="", kratownica="", kuchnia="", prysznic="",
              toaleta="", stolarka="", zaluzje="", rolety="", lazienka="", inne=""):
    return [
        kod, nazwa, typ, stan, ilosc, jm, oddzial, rodzaj, status, termin, plyty,
        statyka, kratownica, kuchnia, prysznic, toaleta, stolarka, zaluzje, rolety, lazienka, inne,
    ]


def _build_workbook(rows: list[list]) -> io.BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lista zasobów"
    ws.append(HEADER)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


@pytest.fixture
def base_rows():
    return [
        _make_row(kod="ACT1", status="Logistyka"),
        _make_row(kod="END1", status="Wysłany do klienta"),
    ]


def test_active_and_ended_statuses(operation_times, base_rows):
    report, records = analyze_workbook(_build_workbook(base_rows))
    by_kod = {r["kod"]: r for r in records}
    assert by_kod["ACT1"]["status_classification"] == "ACTIVE"
    assert by_kod["ACT1"]["is_counted"] is True
    assert by_kod["END1"]["status_classification"] == "ENDED"
    assert by_kod["END1"]["is_counted"] is False


def test_unrecognized_status(operation_times):
    rows = [_make_row(kod="X1", status="Status Nieznany")]
    report, records = analyze_workbook(_build_workbook(rows))
    assert records[0]["status_classification"] == "UNRECOGNIZED"
    assert records[0]["is_counted"] is False
    assert report.unrecognized_status_count == 1


def test_repeated_status_normalized(operation_times):
    rows = [_make_row(kod="X1", status="Logistyka\nLogistyka")]
    report, records = analyze_workbook(_build_workbook(rows))
    assert records[0]["status_classification"] == "ACTIVE"
    assert records[0]["is_counted"] is True


def test_conflict_active_and_ended(operation_times):
    rows = [_make_row(kod="X1", status="Produkcja Zabrze\nOdebrany przez klienta")]
    report, records = analyze_workbook(_build_workbook(rows))
    assert records[0]["status_classification"] == "CONFLICT"
    assert records[0]["is_counted"] is False
    assert report.conflict_count == 1


def test_pavilion_without_module(operation_times):
    rows = [_make_row(kod="X1", nazwa="Pawilon 5x3 nr projektu 1/1/2024", status="Logistyka")]
    _, records = analyze_workbook(_build_workbook(rows))
    assert records[0]["is_counted"] is True


def test_module_1_variants(operation_times):
    rows = [_make_row(kod="X1", nazwa="Pawilon 5x3 (MODUŁ 1) nr projektu 1/1/2024", status="Logistyka")]
    _, records = analyze_workbook(_build_workbook(rows))
    assert records[0]["module_number"] == 1
    assert records[0]["is_counted"] is True


def test_modules_2_plus_skipped(operation_times):
    rows = [_make_row(kod="X1", nazwa="Pawilon 5x3 (moduł 2) nr projektu 1/1/2024", status="Logistyka")]
    report, records = analyze_workbook(_build_workbook(rows))
    assert records[0]["is_counted"] is False
    assert report.module_skipped_count == 1


def test_statyka_hours_multiplied_by_total_project_modules(operation_times):
    rows = [
        _make_row(kod="X1", nazwa="Pawilon 10x3 nr projektu 1/1/2024 (MODUŁ 1)",
                  status="Logistyka", statyka="Tak"),
        _make_row(kod="X2", nazwa="Pawilon 10x3 nr projektu 1/1/2024 (MODUŁ 2)",
                  status="Logistyka", statyka="Tak"),
        _make_row(kod="X3", nazwa="Pawilon 10x3 nr projektu 1/1/2024 (MODUŁ 3)",
                  status="Logistyka", statyka="Tak"),
    ]
    _, records = analyze_workbook(_build_workbook(rows))
    by_kod = {r["kod"]: r for r in records}
    # tylko moduł 1 wchodzi do kolejki, ale za to z godzinami razy 3 moduły
    assert by_kod["X1"]["is_counted"] is True
    assert Decimal(by_kod["X1"]["welding_hours"]) == Decimal("12") * 3
    assert by_kod["X2"]["is_counted"] is False
    assert by_kod["X3"]["is_counted"] is False


def test_single_module_project_not_multiplied(operation_times):
    rows = [_make_row(kod="X1", nazwa="Pawilon 5x3 nr projektu 1/1/2024",
                       status="Logistyka", statyka="Tak")]
    _, records = analyze_workbook(_build_workbook(rows))
    assert Decimal(records[0]["welding_hours"]) == Decimal("12")


def test_empty_value_and_pusty(operation_times):
    rows = [_make_row(kod="X1", status="Logistyka", kuchnia="{pusty}")]
    _, records = analyze_workbook(_build_workbook(rows))
    assert records[0]["kuchnia"] == ""


def test_kuchnia_standard_and_lux(operation_times):
    rows = [
        _make_row(kod="X1", status="Logistyka", kuchnia="Standard"),
        _make_row(kod="X2", status="Logistyka", kuchnia="Lux"),
    ]
    _, records = analyze_workbook(_build_workbook(rows))
    by_kod = {r["kod"]: r for r in records}
    assert Decimal(by_kod["X1"]["hydraulic_hours"]) == Decimal("10")
    assert Decimal(by_kod["X2"]["hydraulic_hours"]) == Decimal("10")


def test_toaleta_variants(operation_times):
    rows = [
        _make_row(kod="X1", status="Logistyka", toaleta="Standard"),
        _make_row(kod="X2", status="Logistyka", toaleta="Komfort"),
        _make_row(kod="X3", status="Logistyka", toaleta="Premium"),
    ]
    _, records = analyze_workbook(_build_workbook(rows))
    by_kod = {r["kod"]: r for r in records}
    # Komfort/Premium mają już wliczone Fibo/Płytki (nie są to osobne dodatki).
    assert Decimal(by_kod["X1"]["hydraulic_hours"]) == Decimal("7")
    assert Decimal(by_kod["X2"]["hydraulic_hours"]) == Decimal("52")
    assert Decimal(by_kod["X3"]["hydraulic_hours"]) == Decimal("55")


def test_toaleta_no_longer_accepts_wc_addon_strings(operation_times):
    # Fibo/Płytki nie są już osobnymi wartościami pola Toaleta (wliczone w
    # Komfort/Premium) — jeśli ktoś wpisze "Fibo" tam, to nierozpoznana
    # wartość, a nie cichy sukces.
    rows = [_make_row(kod="X1", status="Logistyka", toaleta="Fibo")]
    report, records = analyze_workbook(_build_workbook(rows))
    assert records[0]["toaleta"] == ""
    assert Decimal(records[0]["custom_bathroom_hours"]) == Decimal("0")
    assert len(report.unrecognized_values) == 1


def test_lazienka_replaces_toaleta_and_prysznic(operation_times):
    rows = [_make_row(kod="X1", status="Logistyka", lazienka="Standard", toaleta="Premium", prysznic="Tak")]
    _, records = analyze_workbook(_build_workbook(rows))
    assert Decimal(records[0]["hydraulic_hours"]) == Decimal("7")


def test_kuchnia_sums_with_lazienka(operation_times):
    rows = [_make_row(kod="X1", status="Logistyka", kuchnia="Lux", lazienka="Premium")]
    _, records = analyze_workbook(_build_workbook(rows))
    assert Decimal(records[0]["hydraulic_hours"]) == Decimal("10") + Decimal("100")


def test_statyka_and_kratownica_together_is_conflict(operation_times):
    # Pełna konstrukcja/statyka i kratownica wykluczają się wzajemnie.
    rows = [_make_row(kod="X1", status="Logistyka", statyka="Tak", kratownica="Tak")]
    report, records = analyze_workbook(_build_workbook(rows))
    assert records[0]["status_classification"] == "CONFLICT"
    assert records[0]["is_counted"] is False
    assert Decimal(records[0]["welding_hours"]) == Decimal("0")
    assert report.conflict_count == 1


def test_statyka_and_kratownica_together_ignored_for_inactive_status(operation_times):
    # Pawilon, który i tak nigdy nie trafiłby do backlogu (status zakończony),
    # nie powinien zaśmiecać listy konfliktów tą regułą.
    rows = [_make_row(kod="X1", status="Wysłany do klienta", statyka="Tak", kratownica="Tak")]
    report, records = analyze_workbook(_build_workbook(rows))
    assert records[0]["status_classification"] == "ENDED"
    assert records[0]["is_counted"] is False
    assert report.conflict_count == 0


def test_fibo_and_boazeria_sum_in_shared_brigade(operation_times):
    header = HEADER + ["FIBO", "BOAZERIA"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lista zasobów"
    ws.append(header)
    ws.append(_make_row(kod="X1", status="Logistyka") + ["Tak", "Tak"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    report, records = analyze_workbook(buf)
    assert report.fibo_column_present is True
    assert report.boazeria_column_present is True
    assert Decimal(records[0]["fibo_wood_hours"]) == Decimal("50") + Decimal("70")


def test_missing_fibo_boazeria_columns_flagged(operation_times, base_rows):
    report, _ = analyze_workbook(_build_workbook(base_rows))
    assert report.fibo_column_present is False
    assert report.boazeria_column_present is False


def test_parallel_brigades_recorded_independently(operation_times):
    rows = [_make_row(kod="X1", status="Logistyka", kuchnia="Lux", statyka="Tak")]
    _, records = analyze_workbook(_build_workbook(rows))
    assert Decimal(records[0]["hydraulic_hours"]) == Decimal("10")
    assert Decimal(records[0]["welding_hours"]) == Decimal("12")


def test_duplicate_active_kod_flagged_as_conflict(operation_times):
    rows = [
        _make_row(kod="DUP1", status="Logistyka"),
        _make_row(kod="DUP1", status="Logistyka"),
    ]
    report, records = analyze_workbook(_build_workbook(rows))
    assert report.duplicate_count == 1
    counted = [r for r in records if r["kod"] == "DUP1" and r["is_counted"]]
    assert len(counted) == 1


def test_missing_required_columns_raises():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lista zasobów"
    ws.append(["Nazwa"])
    ws.append(["coś"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    with pytest.raises(ImportAnalysisError):
        analyze_workbook(buf)


@pytest.mark.django_db
def test_transactional_import_does_not_change_snapshot_on_failure(operation_times, django_user_model):
    user = django_user_model.objects.create_user(username="admin", password="x")
    rows = [_make_row(kod="ACT1", status="Logistyka")]
    batch1 = ImportBatch.objects.create(file="dummy1.xlsx", original_filename="a.xlsx", uploaded_by=user)
    _, records1 = analyze_workbook(_build_workbook(rows))
    commit_batch(batch1, records1, user)
    assert PavilionSnapshot.objects.filter(import_batch=batch1, is_counted=True).count() == 1

    # symulujemy nieudany drugi import: zły plik (analiza rzuca wyjątek przed zapisem)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Zła nazwa arkusza"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    with pytest.raises(ImportAnalysisError):
        analyze_workbook(buf)

    # poprzednia migawka pozostaje nienaruszona i nadal aktywna
    batch1.refresh_from_db()
    assert batch1.is_active_snapshot is True
    assert PavilionSnapshot.objects.filter(import_batch=batch1, is_counted=True).count() == 1
